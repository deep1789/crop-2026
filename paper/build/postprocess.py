"""Post-process the pandoc output: academic three-line tables, column widths, line numbers, page numbers, figure centring, code-block box."""
import sys, re
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
src, dst = sys.argv[1], sys.argv[2]
d = Document(src)
W = 9071  # usable width in dxa (A4, 2.5 cm margins)

def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")): tblPr.remove(el)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        if edge in ("top", "bottom"):
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "12"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)
    for el in tblPr.findall(qn("w:tblLayout")): tblPr.remove(el)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for el in tblPr.findall(qn("w:tblW")): tblPr.remove(el)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(W)); w.set(qn("w:type"), "dxa"); tblPr.append(w)
    # cell margins
    for el in tblPr.findall(qn("w:tblCellMar")): tblPr.remove(el)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 20), ("left", 60), ("bottom", 20), ("right", 60)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)

def cell_border_bottom(cell, sz=6):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000"); b.append(e); tcPr.append(b)

for tbl in d.tables:
    set_borders(tbl)
    ncol = len(tbl.columns); rows = tbl.rows
    wts = []
    for c in range(ncol):
        L = [len(rows[r].cells[c].text.strip()) for r in range(len(rows))]
        wts.append(min(max(max(L), 6), 48) + 2)
    tot = sum(wts); widths = [int(W * w / tot) for w in wts]
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(w))
    for r in rows:
        for c, cell in enumerate(r.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for el in tcPr.findall(qn("w:tcW")): tcPr.remove(el)
            tw = OxmlElement("w:tcW"); tw.set(qn("w:w"), str(widths[c])); tw.set(qn("w:type"), "dxa"); tcPr.insert(0, tw)
            for p in cell.paragraphs:
                p.paragraph_format.keep_together = True
                for run in p.runs: run.font.size = Pt(8 if ncol > 6 else 8.5)
    for cell in rows[0].cells:
        cell_border_bottom(cell)
        for p in cell.paragraphs:
            for run in p.runs: run.font.bold = True
    trPr = rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for r in rows:   # do not split rows across pages
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)

# highlights: normal-size text
hl = True
for p in d.paragraphs:
    if p.style.name.startswith("Heading") and "Abstract" in p.text: break
    if p.style.name == "Compact":
        for run in p.runs: run.font.size = Pt(10)
# centre paragraphs that contain drawings; keep caption with figure
for p in d.paragraphs:
    if p._p.findall(".//" + qn("w:drawing")):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6)
    if p.style.name == "Source Code":
        pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F2F2F2"); pPr.append(shd)
        for run in p.runs: run.font.size = Pt(8)

# line numbers + page numbers
for sec in d.sections:
    sp = sec._sectPr
    ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); sp.append(ln)
    f = sec.footer; f.is_linked_to_previous = False; p = f.paragraphs[0] if f.paragraphs else f.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t, txt in (("begin", None), ("instr", " PAGE "), ("end", None)):
        r = p.add_run(); r.font.size = Pt(9)
        if t == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._r.append(fc)
for sec in d.sections:
    pm = sec._sectPr.find(qn("w:pgMar"))
    if pm is not None:
        for k, v in (("gutter", "0"), ("header", "708"), ("footer", "708")):
            if pm.get(qn("w:" + k)) is None: pm.set(qn("w:" + k), v)
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
for rp in d.element.iter("{%s}rPr" % M):
    if rp.find("{%s}nor" % M) is not None:
        for st in rp.findall("{%s}sty" % M): rp.remove(st)
# ---- enforce OOXML child order (Word is strict) ----
ORD = {
 "tblPr": ["tblStyle","tblpPr","tblOverlap","bidiVisual","tblStyleRowBandSize","tblStyleColBandSize","tblW","jc","tblCellSpacing","tblInd","tblBorders","shd","tblLayout","tblCellMar","tblLook","tblCaption","tblDescription"],
 "tcPr": ["cnfStyle","tcW","gridSpan","hMerge","vMerge","tcBorders","shd","noWrap","tcMar","textDirection","tcFitText","vAlign","hideMark"],
 "trPr": ["cnfStyle","divId","gridBefore","gridAfter","wBefore","wAfter","cantSplit","trHeight","tblHeader","tblCellSpacing","jc","hidden"],
 "pPr": ["pStyle","keepNext","keepLines","pageBreakBefore","framePr","widowControl","numPr","suppressLineNumbers","pBdr","shd","tabs","suppressAutoHyphens","kinsoku","wordWrap","overflowPunct","topLinePunct","autoSpaceDE","autoSpaceDN","bidi","adjustRightInd","snapToGrid","spacing","ind","contextualSpacing","mirrorIndents","suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap","outlineLvl","divId","cnfStyle","rPr","sectPr","pPrChange"],
 "sectPr": ["headerReference","footerReference","footnotePr","endnotePr","type","pgSz","pgMar","paperSrc","pgBorders","lnNumType","pgNumType","cols","formProt","vAlign","noEndnote","titlePg","textDirection","bidi","rtlGutter","docGrid","printerSettings","sectPrChange"],
}
def reorder(el, order):
    kids = list(el)
    key = lambda k: order.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in order else len(order)
    kids_sorted = sorted(kids, key=key)
    if kids_sorted != kids:
        for k in kids: el.remove(k)
        for k in kids_sorted: el.append(k)
body = d.element
for tag, order in ORD.items():
    for el in body.iter(qn("w:" + tag)): reorder(el, order)
d.core_properties.title = "Duplicated records and region shift in crop-yield benchmarks"; d.core_properties.author = ""
d.save(dst); print("saved", dst)
